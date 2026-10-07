"""Create bounded documents from data, never execute model-generated code."""

import math
import re
from pathlib import Path
from uuid import uuid4
from xml.sax.saxutils import escape

from markdown_it import MarkdownIt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer


def bounded_text(value, limit=40000):
    if not isinstance(value, str) or not value.strip() or len(value) > limit:
        raise ValueError("Нужен непустой текст в пределах лимита.")
    return value


def artifact_path(directory, name, suffix):
    name = re.sub(r"[^\w .-]", "_", bounded_text(name, 100), flags=re.UNICODE)
    name = Path(name).stem.strip(" .")[:60] or "document"
    directory.mkdir(parents=True, exist_ok=True)
    return directory / f"{name}-{uuid4().hex[:8]}{suffix}"


def pdf_font():
    if "Assistant" not in pdfmetrics.getRegisteredFontNames():
        choices = [
            Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
            Path("C:/Windows/Fonts/arial.ttf"),
        ]
        font = next((p for p in choices if p.is_file()), None)
        if not font:
            raise ValueError("На сервере не найден шрифт для русского PDF.")
        pdfmetrics.registerFont(TTFont("Assistant", str(font)))
    return "Assistant"


def make_pdf(directory, name, markdown):
    bounded_text(markdown)
    target = artifact_path(directory, name, ".pdf")
    font = pdf_font()
    body = ParagraphStyle(
        "Body",
        fontName=font,
        fontSize=10,
        leading=16,
        textColor=colors.HexColor("#243247"),
        spaceAfter=8,
    )
    heading = ParagraphStyle(
        "Heading",
        parent=body,
        fontSize=17,
        leading=23,
        spaceBefore=12,
        spaceAfter=10,
        keepWithNext=True,
    )
    code = ParagraphStyle(
        "Code",
        parent=body,
        fontSize=8,
        leading=12,
        backColor=colors.HexColor("#F1F5F9"),
        borderPadding=6,
    )
    story = []
    previous = ""
    for token in MarkdownIt("commonmark", {"html": False}).parse(markdown):
        if token.type == "inline":
            # Escape everything; no remote images, HTML or file inclusion in PDF.
            text = "".join(
                child.content
                for child in token.children or []
                if child.type in {"text", "code_inline", "softbreak", "hardbreak"}
            )
            if text.strip():
                story.append(
                    Paragraph(
                        escape(text).replace("\n", "<br/>"),
                        heading if previous == "heading_open" else body,
                    )
                )
        elif token.type in {"fence", "code_block"}:
            for line in token.content.splitlines():
                story.append(Paragraph(escape(line) or "&#160;", code))
        elif token.type == "hr":
            story.append(Spacer(1, 12))
        previous = token.type
    if not story:
        raise ValueError("В документе нет текста.")

    def footer(canvas, doc):
        canvas.setFont(font, 8)
        canvas.setFillColor(colors.HexColor("#64748B"))
        canvas.drawRightString(550, 28, str(doc.page))

    SimpleDocTemplate(
        str(target), rightMargin=44, leftMargin=44, topMargin=40, bottomMargin=45, title=name
    ).build(story, onFirstPage=footer, onLaterPages=footer)
    return target


def make_excel(directory, name, sheets):
    if not isinstance(sheets, list) or not 1 <= len(sheets) <= 8:
        raise ValueError("В таблице должно быть от 1 до 8 листов.")
    workbook = Workbook()
    workbook.remove(workbook.active)
    total = 0
    for sheet in sheets:
        if not isinstance(sheet, dict):
            raise ValueError("Некорректный лист.")
        title = re.sub(r"[\\/*?:\[\]]", "_", bounded_text(sheet.get("name"), 60))[:31]
        headers, rows = sheet.get("headers"), sheet.get("rows")
        if not isinstance(headers, list) or not 1 <= len(headers) <= 30:
            raise ValueError("Допустимо от 1 до 30 столбцов.")
        if not isinstance(rows, list) or len(rows) > 1000:
            raise ValueError("Допустимо до 1000 строк на лист.")
        total += (len(rows) + 1) * len(headers)
        if total > 15000:
            raise ValueError("Таблица слишком большая: максимум 15000 ячеек.")
        ws = workbook.create_sheet(title)
        for row_num, row in enumerate([headers, *rows], 1):
            if not isinstance(row, list) or len(row) != len(headers):
                raise ValueError("Во всех строках должно быть одинаковое число ячеек.")
            for col, value in enumerate(row, 1):
                if value is not None and not isinstance(value, (str, int, float, bool)):
                    raise ValueError("Ячейки могут содержать текст, числа или пустое значение.")
                if isinstance(value, float) and not math.isfinite(value):
                    raise ValueError("Некорректное число.")
                if isinstance(value, str) and (
                    len(value) > 2000 or re.search(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", value)
                ):
                    raise ValueError(
                        "Текст ячейки слишком длинный или содержит управляющие символы."
                    )
                cell = ws.cell(row_num, col, value)
                if isinstance(value, str):
                    cell.data_type = "s"  # Never interpret untrusted text as formulas.
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if row_num == 1:
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = PatternFill("solid", fgColor="334C73")
                elif row_num % 2 == 0:
                    cell.fill = PatternFill("solid", fgColor="F1F5F9")
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for col in range(1, len(headers) + 1):
            width = max(
                len(str(ws.cell(r, col).value or "")) for r in range(1, min(ws.max_row, 100) + 1)
            )
            ws.column_dimensions[get_column_letter(col)].width = min(50, max(14, width + 3))
    target = artifact_path(directory, name, ".xlsx")
    workbook.save(target)
    return target


def make_text(directory, name, content, format="txt"):
    if format not in {"txt", "md", "csv", "json"}:
        raise ValueError("Доступны TXT, Markdown, CSV и JSON.")
    bounded_text(content)
    if format == "json":
        import json

        json.loads(content)
    target = artifact_path(directory, name, "." + format)
    target.write_text(content, encoding="utf-8-sig" if format == "csv" else "utf-8")
    return target
