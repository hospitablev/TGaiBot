"""Bounded attachment extraction. Called in a disposable worker process."""

import base64
import io
import json
import math
import subprocess
import zipfile
from dataclasses import dataclass, field
from pathlib import Path

from defusedxml import ElementTree
from PIL import Image, ImageOps, UnidentifiedImageError
from pypdf import PdfReader

IMAGE_EXT = {".jpg", ".jpeg", ".png", ".webp"}
VIDEO_EXT = {".mp4", ".mov", ".mkv", ".webm", ".avi", ".m4v"}
AUDIO_EXT = {".mp3", ".wav", ".ogg", ".oga", ".m4a", ".flac", ".opus"}
TEXT_EXT = {
    ".txt",
    ".md",
    ".csv",
    ".tsv",
    ".json",
    ".xml",
    ".yaml",
    ".yml",
    ".py",
    ".js",
    ".ts",
    ".html",
    ".css",
    ".log",
}
SUPPORTED_EXT = IMAGE_EXT | VIDEO_EXT | AUDIO_EXT | TEXT_EXT | {".pdf", ".docx", ".xlsx"}
Image.MAX_IMAGE_PIXELS = 20_000_000


class MediaError(Exception):
    pass


@dataclass
class Prepared:
    text: str = ""
    images: list[str] = field(default_factory=list)
    notes: list[str] = field(default_factory=list)

    def content(self, caption):
        text = caption.strip() or "Проанализируй вложение и кратко объясни его содержание."
        if self.notes:
            text += "\n\nОграничения подготовки: " + " ".join(self.notes)
        if self.text:
            text += "\n\n<attachment_data>\n" + self.text + "\n</attachment_data>"
        return [
            {"type": "text", "text": text},
            *[{"type": "image_url", "image_url": {"url": url}} for url in self.images],
        ]


def image_data(path):
    try:
        with Image.open(path) as original:
            if original.format not in {"JPEG", "PNG", "WEBP"}:
                raise MediaError("Поддерживаются изображения JPEG, PNG и WebP.")
            if getattr(original, "n_frames", 1) > 1:
                raise MediaError("Анимированное изображение отправьте как видео MP4.")
            if original.width * original.height > Image.MAX_IMAGE_PIXELS:
                raise MediaError("Изображение слишком большое: максимум 20 мегапикселей.")
            im = ImageOps.exif_transpose(original)
            im.thumbnail((1280, 1280))
            background = Image.new("RGB", im.size, "white")
            if "A" in im.getbands():
                background.paste(im, mask=im.getchannel("A"))
            else:
                background.paste(im.convert("RGB"))
            output = io.BytesIO()
            background.save(output, format="JPEG", quality=80)
            return "data:image/jpeg;base64," + base64.b64encode(output.getvalue()).decode()
    except (
        UnidentifiedImageError,
        OSError,
        Image.DecompressionBombError,
        Image.DecompressionBombWarning,
    ):
        raise MediaError("Не удалось прочитать изображение или превышен лимит пикселей.")


def limited_text(text, settings, notes=None):
    notes = list(notes or [])
    text = text.strip()
    if not text:
        raise MediaError("В документе нет извлекаемого текста. Отправьте нужные страницы фото.")
    if len(text) > settings.max_document_chars:
        text = text[: settings.max_document_chars]
        notes.append(f"Документ обрезан до первых {settings.max_document_chars} символов.")
    return Prepared(text=text, notes=notes)


def document(path, settings):
    ext = path.suffix.lower()
    if ext == ".xlsx":
        from openpyxl import load_workbook

        try:
            with zipfile.ZipFile(path) as archive:
                entries = archive.infolist()
                if len(entries) > 2000 or sum(e.file_size for e in entries) > 30_000_000:
                    raise MediaError("Excel слишком большой после распаковки (лимит 30 МБ).")
            book = load_workbook(path, read_only=True, data_only=False, keep_links=False)
            try:
                lines, size = [], 0
                for sheet in book.worksheets[:8]:
                    lines.append("Лист: " + sheet.title)
                    for row in sheet.iter_rows(max_row=1000, max_col=30, values_only=True):
                        if not any(value is not None for value in row):
                            continue
                        line = " | ".join(
                            "" if value is None else str(value) for value in row
                        ).rstrip(" |")
                        lines.append(line)
                        size += len(line)
                        if size > settings.max_document_chars:
                            break
                    if size > settings.max_document_chars:
                        break
                return limited_text(
                    "\n".join(lines),
                    settings,
                    [
                        "Excel: только текст и значения первых 8 листов, 1000 строк и 30 столбцов; формулы показаны как текст и не вычисляются. Графики не извлекаются."
                    ],
                )
            finally:
                book.close()
        except MediaError:
            raise
        except Exception:
            raise MediaError("Не удалось прочитать Excel. Пришли файл XLSX без пароля.")
    if ext in TEXT_EXT:
        raw = path.read_bytes()
        try:
            text = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            raise MediaError("Текстовый файл должен быть в UTF-8. Пересохраните файл в UTF-8.")
        if "\x00" in text:
            raise MediaError("Файл содержит двоичные данные, а не поддерживаемый текст UTF-8.")
        return limited_text(text, settings)
    if ext == ".docx":
        try:
            with zipfile.ZipFile(path) as archive:
                entries = archive.infolist()
                if len(entries) > 2000 or sum(e.file_size for e in entries) > 30_000_000:
                    raise MediaError("DOCX слишком большой после распаковки (лимит 30 МБ).")
                root = ElementTree.fromstring(archive.read("word/document.xml"))
                ns = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
                lines = []
                for paragraph in root.iter(ns + "p"):
                    lines.append(
                        "".join(
                            (node.text or "")
                            if node.tag == ns + "t"
                            else "\t"
                            if node.tag == ns + "tab"
                            else "\n"
                            for node in paragraph.iter()
                            if node.tag in {ns + "t", ns + "tab", ns + "br"}
                        )
                    )
            return limited_text(
                "\n".join(lines),
                settings,
                ["DOCX: только текст основного документа и таблиц; без рисунков."],
            )
        except MediaError:
            raise
        except Exception:
            raise MediaError("Не удалось прочитать DOCX. Файл повреждён или защищён.")
    if ext == ".pdf":
        try:
            reader = PdfReader(path)
            if reader.is_encrypted:
                raise MediaError("Зашифрованный PDF не поддерживается. Снимите пароль локально.")
            notes = ["PDF: извлечён только текст; иллюстрации и точная вёрстка не передаются."]
            if len(reader.pages) > settings.max_pdf_pages:
                notes.append(f"Прочитаны только первые {settings.max_pdf_pages} страниц.")
            parts, empty = [], 0
            for index, page in enumerate(reader.pages[: settings.max_pdf_pages]):
                text = page.extract_text() or ""
                if not text.strip():
                    empty += 1
                else:
                    parts.append(f"Страница {index + 1}:\n{text}")
                if sum(map(len, parts)) > settings.max_document_chars:
                    break
            if empty:
                notes.append(f"Страниц без текста: {empty}; OCR не выполнялся.")
            return limited_text("\n\n".join(parts), settings, notes)
        except MediaError:
            raise
        except Exception:
            raise MediaError("Не удалось прочитать PDF. Файл повреждён или слишком сложный.")
    raise MediaError("Формат документа не поддерживается. Используйте PDF, DOCX или текст UTF-8.")


def run_tool(command, timeout=40):
    try:
        result = subprocess.run(
            command,
            capture_output=True,
            timeout=timeout,
            check=False,
            creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
        )
    except FileNotFoundError:
        raise MediaError("Для видео и аудио установите FFmpeg и ffprobe, затем перезапустите бота.")
    except subprocess.TimeoutExpired:
        raise MediaError(
            "Медиафайл обрабатывался слишком долго. Отправьте более короткий фрагмент."
        )
    if result.returncode:
        raise MediaError("Не удалось декодировать медиафайл. Попробуйте MP4 (H.264/AAC).")
    return result.stdout


def probe(path, settings):
    # Restrict external resources in crafted playlists/containers to the local file protocol.
    raw = run_tool(
        [
            settings.ffprobe,
            "-v",
            "error",
            "-protocol_whitelist",
            "file,pipe",
            "-show_format",
            "-show_streams",
            "-of",
            "json",
            str(path),
        ]
    )
    try:
        info = json.loads(raw)
        duration = float(info.get("format", {}).get("duration", 0))
        if not math.isfinite(duration) or duration <= 0:
            raise ValueError
    except (ValueError, TypeError):
        raise MediaError("Не удалось определить длительность файла.")
    if duration > settings.max_video_seconds:
        raise MediaError(f"Максимальная длительность видео/аудио — {settings.max_video_seconds} с.")
    return info, duration


def transcribe(path, settings):
    try:
        from faster_whisper import WhisperModel
    except ImportError:
        raise MediaError('Для распознавания речи установите зависимости: pip install -e ".[audio]"')
    wav = path.parent / "speech.wav"
    run_tool(
        [
            settings.ffmpeg,
            "-nostdin",
            "-v",
            "error",
            "-y",
            "-protocol_whitelist",
            "file,pipe",
            "-i",
            str(path),
            "-vn",
            "-t",
            str(settings.max_video_seconds),
            "-ac",
            "1",
            "-ar",
            "16000",
            str(wav),
        ]
    )
    model = WhisperModel(settings.whisper_model, device="cpu", compute_type="int8")
    segments, _ = model.transcribe(str(wav), beam_size=1, vad_filter=True)
    return "\n".join(f"[{s.start:.1f}–{s.end:.1f} с] {s.text.strip()}" for s in segments)


def audiovisual(path, settings, *, audio_only=False):
    info, duration = probe(path, settings)
    videos = [
        s
        for s in info.get("streams", [])
        if s.get("codec_type") == "video" and not s.get("disposition", {}).get("attached_pic")
    ]
    has_audio = any(s.get("codec_type") == "audio" for s in info.get("streams", []))
    result = Prepared()
    if not audio_only:
        if not videos:
            raise MediaError("В файле нет видеодорожки.")
        stream = videos[0]
        if int(stream.get("width", 0)) * int(stream.get("height", 0)) > 16_000_000:
            raise MediaError("Разрешение видео слишком велико: максимум 16 мегапикселей.")
        count = min(settings.max_frames, max(1, math.ceil(duration / 10)))
        timestamps = [(i + 0.5) * duration / count for i in range(count)]
        for i, timestamp in enumerate(timestamps):
            frame = path.parent / f"frame-{i}.jpg"
            run_tool(
                [
                    settings.ffmpeg,
                    "-nostdin",
                    "-v",
                    "error",
                    "-y",
                    "-protocol_whitelist",
                    "file,pipe",
                    "-ss",
                    str(timestamp),
                    "-i",
                    str(path),
                    "-map",
                    f"0:{stream['index']}",
                    "-frames:v",
                    "1",
                    "-vf",
                    "scale=1280:1280:force_original_aspect_ratio=decrease",
                    "-threads",
                    "1",
                    str(frame),
                ]
            )
            result.images.append(image_data(frame))
        result.text = f"Видео, {duration:.1f} с. Кадры по порядку: " + ", ".join(
            f"{ts:.1f} с" for ts in timestamps
        )
        result.notes.append(
            f"Видео: выборка из {count} кадров, события между кадрами могут быть пропущены."
        )
    if has_audio and settings.transcription:
        try:
            speech = transcribe(path, settings)
            result.text += "\nРасшифровка речи (может содержать ошибки):\n" + (
                speech or "Речь не найдена."
            )
        except Exception:
            if audio_only:
                raise MediaError("Не удалось распознать речь. Проверьте локальную модель Whisper.")
            result.notes.append("Не удалось распознать звук; анализируются только кадры.")
    elif audio_only:
        raise MediaError(
            "Распознавание аудио отключено или в файле нет звука. "
            "Для речи нужен ENABLE_TRANSCRIPTION=true и зависимости audio."
        )
    elif has_audio:
        result.notes.append("Звук не анализировался: локальное распознавание речи отключено.")
    else:
        result.notes.append("Видеофайл не содержит звуковой дорожки.")
    return result


def prepare(path: Path, settings):
    if path.stat().st_size > settings.max_file_bytes:
        raise MediaError("Файл превышает лимит 20 МБ.")
    suffix = path.suffix.lower()
    if suffix in IMAGE_EXT:
        return Prepared(images=[image_data(path)])
    if suffix in VIDEO_EXT:
        return audiovisual(path, settings)
    if suffix in AUDIO_EXT:
        return audiovisual(path, settings, audio_only=True)
    return document(path, settings)
